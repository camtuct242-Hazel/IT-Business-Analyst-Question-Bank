# bank_builder/exporter.py
import json
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from blueprint_data import BLUEPRINT
from validator import ALL_EXAMS

OUTPUT_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
XLSX_PATH = os.path.join(OUTPUT_DIR, "ITBA_Junior_Question_Bank.xlsx")
JSON_PATH = os.path.join(OUTPUT_DIR, "ITBA_Junior_Question_Bank.json")
SITE_JSON_PATH = os.path.join(OUTPUT_DIR, "site", "ITBA_Junior_Question_Bank.json")
DOCX_PATH = os.path.join(OUTPUT_DIR, "ITBA_Junior_Practice_Guide.docx")

# Style definitions for Excel
HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
ZEBRA_FILL = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
THIN_BORDER = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

def style_worksheet(ws, min_col=1, max_col=10, min_row=1, max_row=1):
    ws.views.sheetView[0].showGridLines = True
    # Freeze top row
    ws.freeze_panes = "A2"
    
    # Header row
    for col in range(min_col, max_col + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER
    ws.row_dimensions[1].height = 28
    
    # Data rows
    for r in range(2, max_row + 1):
        is_even = (r % 2 == 0)
        for col in range(min_col, max_col + 1):
            cell = ws.cell(row=r, column=col)
            cell.border = THIN_BORDER
            cell.font = Font(name="Calibri", size=10)
            if is_even:
                cell.fill = ZEBRA_FILL
            # Alignment heuristics
            if col in [1, 2, 3, 6, 7]: # IDs, numbers, difficulties
                cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
            else:
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                
    # Auto adjust column widths with reasonable caps
    for col in range(min_col, max_col + 1):
        col_letter = get_column_letter(col)
        max_len = 0
        for r in range(1, min(max_row + 1, 25)): # sample first 25 rows
            val = str(ws.cell(row=r, column=col).value or "")
            max_len = max(max_len, len(val))
        calc_width = min(max(max_len + 3, 12), 45)
        ws.column_dimensions[col_letter].width = calc_width

def export_excel(all_questions):
    print("Generating Excel file: ITBA_Junior_Question_Bank.xlsx ...")
    wb = openpyxl.Workbook()
    # remove default sheet
    wb.remove(wb.active)
    
    # 1. Sheet Blueprint
    ws_bp = wb.create_sheet(title="Blueprint")
    bp_headers = [
        "Exam ID", "Exam Title", "Topic", "Learning Objectives",
        "Easy Count", "Intermediate Count", "Advanced Count", "Total Questions"
    ]
    ws_bp.append(bp_headers)
    for b in BLUEPRINT:
        ws_bp.append([
            b["exam_id"], b["exam_title"], b["topic"], b["learning_objectives"],
            b["easy_count"], b["intermediate_count"], b["advanced_count"], b["total_questions"]
        ])
    style_worksheet(ws_bp, min_col=1, max_col=len(bp_headers), min_row=1, max_row=len(BLUEPRINT)+1)
    ws_bp.column_dimensions["A"].width = 12
    ws_bp.column_dimensions["B"].width = 35
    ws_bp.column_dimensions["C"].width = 35
    ws_bp.column_dimensions["D"].width = 50
    ws_bp.column_dimensions["E"].width = 14
    ws_bp.column_dimensions["F"].width = 18
    ws_bp.column_dimensions["G"].width = 16
    ws_bp.column_dimensions["H"].width = 16

    # 2. Sheet Question Paper
    ws_qp = wb.create_sheet(title="Question Paper")
    qp_headers = [
        "Question ID", "Exam ID", "Question Number", "Topic", "Subtopic",
        "Difficulty", "Question Type", "Scenario", "Question Text",
        "Option A", "Option B", "Option C", "Option D"
    ]
    ws_qp.append(qp_headers)
    for q in all_questions:
        ws_qp.append([
            q["question_id"], q["exam_id"], q["question_number"], q["topic"], q["subtopic"],
            q["difficulty"], q["question_type"], q["scenario"] or "", q["question_text"],
            q["option_a"], q["option_b"], q["option_c"], q["option_d"]
        ])
    style_worksheet(ws_qp, min_col=1, max_col=len(qp_headers), min_row=1, max_row=len(all_questions)+1)
    ws_qp.column_dimensions["A"].width = 14
    ws_qp.column_dimensions["B"].width = 12
    ws_qp.column_dimensions["C"].width = 16
    ws_qp.column_dimensions["D"].width = 30
    ws_qp.column_dimensions["E"].width = 25
    ws_qp.column_dimensions["F"].width = 12
    ws_qp.column_dimensions["G"].width = 20
    ws_qp.column_dimensions["H"].width = 45
    ws_qp.column_dimensions["I"].width = 55
    ws_qp.column_dimensions["J"].width = 35
    ws_qp.column_dimensions["K"].width = 35
    ws_qp.column_dimensions["L"].width = 35
    ws_qp.column_dimensions["M"].width = 35

    # 3. Sheet Answer Key
    ws_ak = wb.create_sheet(title="Answer Key")
    ak_headers = [
        "Question ID", "Correct Answer", "Explanation",
        "Why A Is Incorrect", "Why B Is Incorrect", "Why C Is Incorrect", "Why D Is Incorrect",
        "Key Concept", "Learning Point", "Reference URL"
    ]
    ws_ak.append(ak_headers)
    for q in all_questions:
        ws_ak.append([
            q["question_id"], q["correct_answer"], q["explanation"],
            q["why_a_incorrect"], q["why_b_incorrect"], q["why_c_incorrect"], q["why_d_incorrect"],
            q["key_concept"], q["learning_point"], q["reference_url"]
        ])
    style_worksheet(ws_ak, min_col=1, max_col=len(ak_headers), min_row=1, max_row=len(all_questions)+1)
    ws_ak.column_dimensions["A"].width = 14
    ws_ak.column_dimensions["B"].width = 15
    ws_ak.column_dimensions["C"].width = 55
    ws_ak.column_dimensions["D"].width = 35
    ws_ak.column_dimensions["E"].width = 35
    ws_ak.column_dimensions["F"].width = 35
    ws_ak.column_dimensions["G"].width = 35
    ws_ak.column_dimensions["H"].width = 30
    ws_ak.column_dimensions["I"].width = 45
    ws_ak.column_dimensions["J"].width = 40

    # 4. Sheet Progress Tracker
    ws_pt = wb.create_sheet(title="Progress Tracker")
    pt_headers = [
        "Exam ID", "Exam Title", "Target Score (out of 20)",
        "Actual Score", "Accuracy (%)", "Weak Topics", "Review Notes", "Retake Score"
    ]
    ws_pt.append(pt_headers)
    for b in BLUEPRINT:
        ws_pt.append([
            b["exam_id"], b["exam_title"], 20,
            "", "", "", "", ""
        ])
    # Add formulas for accuracy: e.g. =IF(D2="","",D2/C2*100)
    for r in range(2, len(BLUEPRINT) + 2):
        ws_pt.cell(row=r, column=5).value = f'=IF(D{r}="","",D{r}/C{r}*100)'
    style_worksheet(ws_pt, min_col=1, max_col=len(pt_headers), min_row=1, max_row=len(BLUEPRINT)+1)
    ws_pt.column_dimensions["A"].width = 12
    ws_pt.column_dimensions["B"].width = 35
    ws_pt.column_dimensions["C"].width = 22
    ws_pt.column_dimensions["D"].width = 14
    ws_pt.column_dimensions["E"].width = 16
    ws_pt.column_dimensions["F"].width = 30
    ws_pt.column_dimensions["G"].width = 40
    ws_pt.column_dimensions["H"].width = 14

    wb.save(XLSX_PATH)
    print(f"Excel file saved successfully to: {XLSX_PATH}")

def export_json(all_questions):
    print("Generating JSON file: ITBA_Junior_Question_Bank.json ...")
    data = {
        "title": "IT Business Analyst Question Bank (200 Questions / 10 Exams)",
        "target_audience": "BA Intern, BA Fresher, IT BA Junior",
        "total_questions": len(all_questions),
        "total_exams": len(BLUEPRINT),
        "blueprint": BLUEPRINT,
        "questions": all_questions
    }
    with open(JSON_PATH, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    print(f"JSON file saved successfully to: {JSON_PATH}")
    os.makedirs(os.path.dirname(SITE_JSON_PATH), exist_ok=True)
    with open(SITE_JSON_PATH, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    print(f"Site JSON file saved successfully to: {SITE_JSON_PATH}")

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def export_docx(all_questions):
    print("Generating Word Document: ITBA_Junior_Practice_Guide.docx ...")
    doc = docx.Document()
    
    # Set page margins: 0.75 in
    for s in doc.sections:
        s.top_margin = Inches(0.75)
        s.bottom_margin = Inches(0.75)
        s.left_margin = Inches(0.75)
        s.right_margin = Inches(0.75)
        
    # Styles
    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("IT BUSINESS ANALYST QUESTION BANK\n(200 CÂU HỎI / 10 ĐỀ THI ĐỘC LẬP)")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_sub.add_run("Bộ tài liệu ôn luyện & Đánh giá năng lực toàn diện dành cho BA Intern, BA Fresher & IT BA Junior\nChuẩn IIBA BABOK v3, Scrum Guide 2020, BPMN 2.0, UML 2.5, SQL & REST API")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(11)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
    
    doc.add_paragraph() # spacer
    
    # Overview Box
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    # Light blue background
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'F2F5F9')
    tcPr.append(shd)
    
    p_box = cell.paragraphs[0]
    r_box_title = p_box.add_run("HƯỚNG DẪN LUYỆN TẬP VÀ QUY CÁCH ĐỀ THI\n")
    r_box_title.bold = True
    r_box_title.font.size = Pt(11)
    r_box_title.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    r_box_body = p_box.add_run(
        "• Cấu trúc: 10 Đề thi độc lập (EXAM01 – EXAM10), mỗi đề chính xác 20 câu hỏi trắc nghiệm 4 lựa chọn (A, B, C, D).\n"
        "• Phân bổ độ khó mỗi đề: 10 câu Dễ (50%), 5 câu Khá (25%), 5 câu Khó (25%).\n"
        "• Thời gian mục tiêu: 20 phút / đề (1 phút / câu). Điểm tối đa: 20 điểm (1 điểm / câu).\n"
        "• Cấu trúc tài liệu:\n"
        "   + PHẦN A (Trang sau): ĐỀ THI KHÔNG ĐÁP ÁN — Dùng để làm bài tự luyện tập nghiêm túc bấm giờ.\n"
        "   + PHẦN B: ĐÁP ÁN VÀ GIẢI THÍCH CHI TIẾT — Đặt sau toàn bộ 10 đề thi, giải thích lý do đúng/sai, kiến thức trọng tâm và tài liệu tham khảo chuẩn mực."
    )
    r_box_body.font.size = Pt(10)
    
    doc.add_page_break()
    
    # ==========================================
    # PART A: QUESTION PAPERS
    # ==========================================
    h_part_a = doc.add_heading(level=1)
    r_pa = h_part_a.add_run("PHẦN A — ĐỀ THI THỰC HÀNH (QUESTION PAPERS)")
    r_pa.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    r_pa.font.bold = True
    
    p_pa_note = doc.add_paragraph()
    r_pan = p_pa_note.add_run("Lưu ý: Phần A hoàn toàn KHÔNG hiển thị đáp án để người học tự luyện tập khách quan. Hãy ghi câu trả lời ra giấy và kiểm tra với Phần B sau khi hoàn tất.")
    r_pan.font.italic = True
    r_pan.font.size = Pt(10)
    
    # Group questions by exam
    exams_dict = {}
    for q in all_questions:
        eid = q["exam_id"]
        if eid not in exams_dict:
            exams_dict[eid] = []
        exams_dict[eid].append(q)
        
    for b in BLUEPRINT:
        eid = b["exam_id"]
        q_list = exams_dict[eid]
        
        doc.add_heading(f"{b['exam_title']}", level=2)
        p_desc = doc.add_paragraph()
        r_desc = p_desc.add_run(f"Mục tiêu kiến thức: {b['learning_objectives']}\nQuy cách: 20 câu (10 Dễ, 5 Khá, 5 Khó) | Thời gian làm bài: 20 phút | Thang điểm: 20 điểm.")
        r_desc.font.size = Pt(9.5)
        r_desc.font.italic = True
        r_desc.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
        
        for q in q_list:
            p_q = doc.add_paragraph()
            p_q.paragraph_format.space_before = Pt(8)
            p_q.paragraph_format.space_after = Pt(2)
            
            # Question Header: [Q001] Câu 1 (Độ khó: Dễ | Chủ đề: ...)
            r_num = p_q.add_run(f"[{q['question_id']}] Câu {q['question_number']} ")
            r_num.bold = True
            r_num.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
            
            r_tag = p_q.add_run(f"[{q['difficulty']} • {q['subtopic']}]\n")
            r_tag.font.size = Pt(9)
            r_tag.font.color.rgb = RGBColor(0x70, 0x30, 0xA0) if q['difficulty'] == 'Khó' else (RGBColor(0xED, 0x7D, 0x31) if q['difficulty'] == 'Khá' else RGBColor(0x38, 0x57, 0x23))
            
            # Scenario if present
            if q.get("scenario"):
                p_sc = doc.add_paragraph()
                p_sc.paragraph_format.left_indent = Inches(0.2)
                p_sc.paragraph_format.space_after = Pt(2)
                r_sc_label = p_sc.add_run("Tình huống: ")
                r_sc_label.bold = True
                r_sc_label.font.size = Pt(9.5)
                r_sc = p_sc.add_run(q["scenario"])
                r_sc.font.italic = True
                r_sc.font.size = Pt(9.5)
                r_sc.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
                
            p_qt = doc.add_paragraph()
            p_qt.paragraph_format.space_after = Pt(4)
            r_qt = p_qt.add_run(q["question_text"])
            r_qt.bold = True
            r_qt.font.size = Pt(10)
            
            # Options A, B, C, D
            for opt_key, opt_label in [("option_a", "A"), ("option_b", "B"), ("option_c", "C"), ("option_d", "D")]:
                p_opt = doc.add_paragraph()
                p_opt.paragraph_format.left_indent = Inches(0.25)
                p_opt.paragraph_format.space_after = Pt(1)
                r_opt_lbl = p_opt.add_run(f"{opt_label}. ")
                r_opt_lbl.bold = True
                r_opt_lbl.font.size = Pt(9.5)
                r_opt_txt = p_opt.add_run(q[opt_key])
                r_opt_txt.font.size = Pt(9.5)
                
        doc.add_page_break()

    # ==========================================
    # PART B: ANSWER KEYS AND EXPLANATIONS
    # ==========================================
    h_part_b = doc.add_heading(level=1)
    r_pb = h_part_b.add_run("PHẦN B — ĐÁP ÁN VÀ GIẢI THÍCH CHI TIẾT (ANSWER KEYS & EXPLANATIONS)")
    r_pb.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    r_pb.font.bold = True
    
    p_pb_note = doc.add_paragraph()
    r_pbn = p_pb_note.add_run("Bao gồm đáp án chính xác, phân tích vì sao đáp án đúng, phân tích chi tiết các đáp án nhiễu (distractors), kiến thức trọng tâm và nguồn tài liệu đối soát.")
    r_pbn.font.italic = True
    r_pbn.font.size = Pt(10)

    for b in BLUEPRINT:
        eid = b["exam_id"]
        q_list = exams_dict[eid]
        
        doc.add_heading(f"ĐÁP ÁN & GIẢI THÍCH {b['exam_title'].upper()}", level=2)
        
        for q in q_list:
            p_ans_head = doc.add_paragraph()
            p_ans_head.paragraph_format.space_before = Pt(8)
            p_ans_head.paragraph_format.space_after = Pt(2)
            
            r_aq = p_ans_head.add_run(f"[{q['question_id']}] Câu {q['question_number']}: ")
            r_aq.bold = True
            r_aq.font.size = Pt(10)
            r_aq.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
            
            r_ac = p_ans_head.add_run(f"ĐÁP ÁN ĐÚNG: {q['correct_answer']}\n")
            r_ac.bold = True
            r_ac.font.size = Pt(10.5)
            r_ac.font.color.rgb = RGBColor(0x00, 0x80, 0x00)
            
            # Details block
            p_det = doc.add_paragraph()
            p_det.paragraph_format.left_indent = Inches(0.2)
            p_det.paragraph_format.space_after = Pt(4)
            
            r_exp_lbl = p_det.add_run("• Giải thích lý do đúng: ")
            r_exp_lbl.bold = True
            r_exp_lbl.font.size = Pt(9.5)
            r_exp = p_det.add_run(f"{q['explanation']}\n")
            r_exp.font.size = Pt(9.5)
            
            r_why_lbl = p_det.add_run("• Phân tích đáp án sai:\n")
            r_why_lbl.bold = True
            r_why_lbl.font.size = Pt(9.5)
            for opt_key, opt_label in [("why_a_incorrect", "A"), ("why_b_incorrect", "B"), ("why_c_incorrect", "C"), ("why_d_incorrect", "D")]:
                if opt_label != q["correct_answer"]:
                    p_det.add_run(f"   - Lựa chọn {opt_label}: {q[opt_key]}\n").font.size = Pt(9)
                    
            r_kc_lbl = p_det.add_run("• Kiến thức trọng tâm: ")
            r_kc_lbl.bold = True
            r_kc_lbl.font.size = Pt(9.5)
            p_det.add_run(f"{q['key_concept']}\n").font.size = Pt(9.5)
            
            r_lp_lbl = p_det.add_run("• Bài học ghi nhớ: ")
            r_lp_lbl.bold = True
            r_lp_lbl.font.size = Pt(9.5)
            p_det.add_run(f"{q['learning_point']}\n").font.size = Pt(9.5)
            
            r_ref_lbl = p_det.add_run("• Tài liệu tham khảo: ")
            r_ref_lbl.bold = True
            r_ref_lbl.font.size = Pt(9.5)
            r_ref = p_det.add_run(f"{q['reference_url']}\n")
            r_ref.font.size = Pt(9)
            r_ref.font.italic = True
            r_ref.font.color.rgb = RGBColor(0x00, 0x45, 0x86)
            
        doc.add_page_break()
        
    doc.save(DOCX_PATH)
    print(f"Word Document saved successfully to: {DOCX_PATH}")

def main():
    from validator import validate_all
    all_questions = validate_all()
    export_excel(all_questions)
    export_json(all_questions)
    export_docx(all_questions)
    print("\nALL FILES CREATED SUCCESSFULLY!")

if __name__ == "__main__":
    main()
